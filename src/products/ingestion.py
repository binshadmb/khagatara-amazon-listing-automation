"""Deterministic extraction of product source files into text."""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".json",
    ".csv",
    ".xlsx",
    ".pdf",
}


class IngestionResult:
    def __init__(
        self,
        *,
        source_type: str,
        text: str,
        metadata: dict[str, Any],
    ) -> None:
        self.source_type = source_type
        self.text = text
        self.metadata = metadata

    def to_dict(self) -> dict[str, Any]:
        return {
            "source_type": self.source_type,
            "text": self.text,
            "metadata": self.metadata,
        }


def ingest_file(path: str | Path) -> IngestionResult:
    """Extract deterministic text/data from a supported source file."""

    source = Path(path)

    if not source.exists():
        raise FileNotFoundError(f"Source file not found: {source}")

    extension = source.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported source type: {extension or '(none)'}"
        )

    if extension == ".txt":
        return _ingest_text(source)

    if extension == ".json":
        return _ingest_json(source)

    if extension == ".csv":
        return _ingest_csv(source)

    if extension == ".xlsx":
        return _ingest_xlsx(source)

    if extension == ".pdf":
        return _ingest_pdf(source)

    raise ValueError(f"Unsupported source type: {extension}")


def _metadata(path: Path, source_type: str) -> dict[str, Any]:
    return {
        "filename": path.name,
        "size": path.stat().st_size,
        "source_type": source_type,
    }


def _ingest_text(path: Path) -> IngestionResult:
    text = path.read_text(encoding="utf-8", errors="replace").strip()

    return IngestionResult(
        source_type="txt",
        text=text,
        metadata=_metadata(path, "txt"),
    )


def _ingest_json(path: Path) -> IngestionResult:
    with path.open("r", encoding="utf-8") as file:
        data = json.load(file)

    text = _structured_to_text(data)

    return IngestionResult(
        source_type="json",
        text=text,
        metadata=_metadata(path, "json"),
    )


def _ingest_csv(path: Path) -> IngestionResult:
    rows: list[list[str]] = []

    with path.open(
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline="",
    ) as file:
        reader = csv.reader(file)

        for index, row in enumerate(reader):
            if index >= 500:
                break

            rows.append([cell.strip() for cell in row])

    text = "\n".join(
        " | ".join(row)
        for row in rows
        if any(row)
    )

    return IngestionResult(
        source_type="csv",
        text=text,
        metadata={
            **_metadata(path, "csv"),
            "rows": len(rows),
        },
    )


def _ingest_xlsx(path: Path) -> IngestionResult:
    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    sheet_names = list(workbook.sheetnames)
    lines: list[str] = []
    row_count = 0

    try:
        for worksheet in workbook.worksheets:
            lines.append(f"[Sheet: {worksheet.title}]")

            for row_index, row in enumerate(
                worksheet.iter_rows(values_only=True)
            ):
                if row_index >= 500:
                    break

                values = [
                    str(value).strip()
                    for value in row
                    if value is not None
                ]

                if values:
                    lines.append(" | ".join(values))
                    row_count += 1
    finally:
        workbook.close()

    return IngestionResult(
        source_type="xlsx",
        text="\n".join(lines),
        metadata={
            **_metadata(path, "xlsx"),
            "sheets": sheet_names,
            "rows": row_count,
        },
    )


def _ingest_pdf(path: Path) -> IngestionResult:
    reader = PdfReader(str(path))

    pages: list[str] = []

    for page in reader.pages:
        text = page.extract_text() or ""
        text = text.strip()

        if text:
            pages.append(text)

    text = "\n\n".join(pages)

    return IngestionResult(
        source_type="pdf",
        text=text,
        metadata={
            **_metadata(path, "pdf"),
            "pages": len(reader.pages),
            "text_pages": len(pages),
            "text_extracted": bool(text),
        },
    )


def _structured_to_text(value: Any, prefix: str = "") -> str:
    lines: list[str] = []

    if isinstance(value, dict):
        for key, item in value.items():
            label = f"{prefix}{key}"

            if isinstance(item, (dict, list)):
                nested = _structured_to_text(item, f"{label}.")
                if nested:
                    lines.append(nested)
            else:
                lines.append(f"{label}: {item}")

    elif isinstance(value, list):
        for index, item in enumerate(value):
            nested = _structured_to_text(
                item,
                f"{prefix}{index}.",
            )

            if nested:
                lines.append(nested)

    else:
        lines.append(f"{prefix}{value}")

    return "\n".join(lines)
